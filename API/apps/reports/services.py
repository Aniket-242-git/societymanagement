"""Dashboard aggregation + Excel export helpers (openpyxl)."""
import datetime as dt
from io import BytesIO

from django.db.models import Count, Sum
from openpyxl import Workbook

from API.apps.expenses.models import Expense
from API.apps.flats.models import Flat
from API.apps.issues.models import Issue
from API.apps.payments.models import MaintenancePayment


def month_bounds(year, month):
    start = dt.date(int(year), int(month), 1)
    end = dt.date(int(year) + (1 if month == 12 else 0), 1 if month == 12 else month + 1, 1)
    return start, end


def collection_summary(year=None, month=None):
    """Approved collections vs pending for a period (defaults: all time)."""
    approved = MaintenancePayment.objects.filter(is_deleted=False, status="approved")
    pending = MaintenancePayment.objects.filter(is_deleted=False, status="pending")
    if year and month:
        s, e = month_bounds(year, month)
        approved = approved.filter(approved_at__date__gte=s, approved_at__date__lt=e)
        pending = pending.filter(created_at__date__gte=s, created_at__date__lt=e)
    elif year:
        approved = approved.filter(approved_at__year=int(year))
        pending = pending.filter(created_at__year=int(year))
    a_sum = approved.aggregate(total=Sum("amount"))["total"] or 0
    p_sum = pending.aggregate(total=Sum("amount"))["total"] or 0
    return {
        "approved_total": float(a_sum),
        "pending_total": float(p_sum),
        "approved_count": approved.count(),
        "pending_count": pending.count(),
    }


def flat_wise_due_list(year, month):
    """Active flats with no approved payment in the given month => due list."""
    s, e = month_bounds(year, month)
    paid_flat_ids = set(
        MaintenancePayment.objects.filter(
            is_deleted=False, status="approved", approved_at__date__gte=s, approved_at__date__lt=e
        ).values_list("flat_id", flat=True)
    )
    flats = Flat.objects.filter(is_active=True).select_related("wing").order_by("wing__name", "flat_no")
    rows = []
    for f in flats:
        if f.id not in paid_flat_ids:
            rows.append({
                "flat": str(f), "wing": f.wing.name, "owner": f.owner_name,
                "monthly_maintenance": float(f.monthly_maintenance),
                "period": f"{month:02d}/{year}",
            })
    return rows


def expense_breakdown(year=None, month=None):
    qs = Expense.objects.filter(is_active=True)
    if year and month:
        s, e = month_bounds(year, month)
        qs = qs.filter(expense_date__gte=s, expense_date__lt=e)
    elif year:
        qs = qs.filter(expense_date__year=int(year))
    rows = list(qs.values("category").annotate(total=Sum("amount"), count=Count("id")).order_by("-total"))
    return [{"category": r["category"], "total": float(r["total"]), "count": r["count"]} for r in rows]


def income_vs_expense(year=None, month=None):
    inc = collection_summary(year, month)["approved_total"]
    exp_rows = expense_breakdown(year, month)
    exp = sum(r["total"] for r in exp_rows)
    return {"income": inc, "expense": exp, "balance": inc - exp, "by_category": exp_rows}


def dashboard_stats(user):
    """Aggregated counters for admin/resident dashboard cards."""
    now = dt.date.today()
    data = {
        "active_flats": Flat.objects.filter(is_active=True).count(),
        "pending_payments": MaintenancePayment.objects.filter(is_deleted=False, status="pending").count(),
        "open_issues": Issue.objects.filter(is_active=True).exclude(status="resolved").count(),
        "announcements": 0,
        "this_month_collection": collection_summary(now.year, now.month)["approved_total"],
        "this_month_expense": sum(r["total"] for r in expense_breakdown(now.year, now.month)),
    }
    if user.is_staff_role:
        from API.apps.announcements.models import Announcement
        data["announcements"] = Announcement.objects.filter(is_active=True).count()
    else:
        my_payments = MaintenancePayment.objects.filter(is_deleted=False, flat__owner=user)
        data["my_pending"] = my_payments.filter(status="pending").count()
        data["my_approved_this_year"] = float(
            my_payments.filter(status="approved", approved_at__year=now.year)
            .aggregate(t=Sum("amount"))["t"] or 0
        )
    return data


# ------------------------------------------------------------------ Excel export
def _write_sheet(ws, headers, rows):
    ws.append(headers)
    try:
        from openpyxl.styles import Font
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)
    except Exception:  # pragma: no cover
        pass
    for row in rows:
        ws.append(row)
    for col_idx in range(1, len(headers) + 1):
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = 24


def _xlsx_response_bytes(wb):
    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio


def export_collections_xlsx(year=None, month=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Collections"
    qs = MaintenancePayment.objects.filter(is_deleted=False, status="approved").select_related("flat", "submitted_by")
    if year and month:
        s, e = month_bounds(year, month)
        qs = qs.filter(approved_at__date__gte=s, approved_at__date__lt=e)
    rows = [
        [p.flat.__str__(), p.submitted_by.username if p.submitted_by else "",
         float(p.amount), p.receipt_no, p.get_mode_display(),
         p.approved_at.strftime("%Y-%m-%d") if p.approved_at else ""]
        for p in qs.order_by("-created_at")
    ]
    _write_sheet(ws, ["Flat", "Submitted By", "Amount", "Receipt No", "Mode", "Approved On"], rows)
    return _xlsx_response_bytes(wb)


def export_expenses_xlsx(year=None, month=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Expenses"
    qs = Expense.objects.filter(is_active=True)
    if year and month:
        s, e = month_bounds(year, month)
        qs = qs.filter(expense_date__gte=s, expense_date__lt=e)
    rows = [[x.title, x.get_category_display(), float(x.amount), x.paid_to, x.expense_date.strftime("%Y-%m-%d")]
            for x in qs.order_by("-expense_date")]
    _write_sheet(ws, ["Title", "Category", "Amount", "Paid To", "Date"], rows)
    return _xlsx_response_bytes(wb)


def export_due_list_xlsx(year, month):
    wb = Workbook()
    ws = wb.active
    ws.title = "Due List"
    rows = [[r["flat"], r["owner"], r["monthly_maintenance"], r["period"]] for r in flat_wise_due_list(year, month)]
    _write_sheet(ws, ["Flat", "Owner", "Monthly Maintenance", "Period"], rows)
    return _xlsx_response_bytes(wb)
